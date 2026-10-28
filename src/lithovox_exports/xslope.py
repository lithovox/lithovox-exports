"""Een dwarsdoorsnede als invoerbestand voor XSLOPE (https://github.com/njones61/xslope):
een kopie van het lege sjabloon van XSLOPE (resources/) met de materialen
en de grondpolygonen ingevuld, en de kleuren van de materialen in
`<naam>_styles.json` ernaast (waar XSLOPE Studio ze leest)."""

import json
import re
import tempfile
import zipfile
from pathlib import Path

import openpyxl
from geolib.models.dstability.internal import PersistableSoil
from openpyxl.utils import get_column_letter

from .dstability import library_soils
from .profile import ProfileSlice, SoilPolygon

TEMPLATE = Path(__file__).parent / "resources" / "xslope_input_template.xlsx"
DECIMALS = 3
GAMMA_WATER = 9.81
MOHR_COULOMB = {
    "MohrCoulombAdvanced": "MohrCoulombAdvancedShearStrengthModel",
    "MohrCoulombClassic": "MohrCoulombClassicShearStrengthModel",
}
# De rijen van een polygoonblok op het blad 'polygon' van het sjabloon.
HEADER_ROW, TYPE_ROW, MAT_ID_ROW, NAME_ROW, SIZE_ROW, XY_ROW, FIRST_VERTEX_ROW = 4, 5, 6, 7, 8, 9, 10


def _properties(soil: PersistableSoil) -> dict[str, object]:
    """De kolommen van het blad 'mat' voor een D-Stability-grondsoort (0 betekent daar: niet ingevuld)."""
    props: dict[str, object] = {}
    above, below = soil.VolumetricWeightAbovePhreaticLevel, soil.VolumetricWeightBelowPhreaticLevel
    if above > 0:
        props["g"] = above
    # XSLOPE weigert een verzadigd gewicht onder het droge.
    if below > 0 and below >= above:
        props["gsat"] = below
    # XSLOPE kent één sterkte per materiaal: die van D-Stability onder de freatische lijn.
    model_type = soil.ShearStrengthModelTypeBelowPhreaticLevel
    model = MOHR_COULOMB.get(getattr(model_type, "value", model_type))
    if model:
        strength = getattr(soil, model)
        props.update(option="mc", c=strength.Cohesion, f=strength.FrictionAngle)
    return props


def _materials(section: ProfileSlice, polygons: list[SoilPolygon], library: dict | None) -> tuple[list[tuple[str, dict, str | None]], dict[int, int]]:
    """De materialen (naam, kolommen van 'mat' en de modelkleur van de eerste
    grondsoort die het materiaal heeft) en het materiaalnummer (vanaf 1) van
    elke grondsoort."""
    soils = [PersistableSoil(**data) for data in (library or {}).get("soils") or []]
    find = library_soils(soils, (library or {}).get("mapping"))
    materials: list[tuple[str, dict, str | None]] = []
    ids: dict[object, int] = {}
    material_of: dict[int, int] = {}
    for soil in sorted({p.soil for p in polygons}):
        name = section.soils[soil]
        stix = find(name)
        key = ("stix", id(stix)) if stix is not None else ("model", soil)
        if key not in ids:
            color = section.colors.get(name)
            materials.append((stix.Name or stix.Code, _properties(stix), color) if stix is not None else (name, {}, color))
            ids[key] = len(materials)
        material_of[soil] = ids[key]
    return materials, material_of


def _styles(materials: list[tuple[str, dict, str | None]]) -> dict:
    """De stijlen van XSLOPE Studio (zijn `<bestand>_styles.json`): elk
    materiaal in zijn modelkleur, ondoorzichtig (XSLOPE vult standaard met
    60 % dekking)."""
    return {
        "version": 1,
        "materials": {
            # Materiaalnummer vanaf 0.
            str(i): {"color": color.lower(), "alpha": 1.0}
            for i, (_, _, color) in enumerate(materials)
            if color and re.fullmatch(r"#[0-9a-fA-F]{6}", color)
        },
    }


def _points(polygon: SoilPolygon) -> list[tuple[float, float]]:
    """De hoekpunten van de polygoon, afgerond, zonder herhalingen (XSLOPE sluit hem zelf)."""
    points: list[tuple[float, float]] = []
    for x, z in polygon.points:
        point = (round(x, DECIMALS), round(z, DECIMALS))
        if not points or point != points[-1]:
            points.append(point)
    if len(points) > 1 and points[0] == points[-1]:
        points.pop()
    return points


def _add_block(ws, n: int) -> None:
    """Zet de kopjes van een polygoonblok voorbij de blokken van het sjabloon (zoals het sjabloon dat doet)."""
    x, y = 3 * n + 1, 3 * n + 2
    col = get_column_letter(y)
    ws.cell(TYPE_ROW, x, "Type:")
    ws.cell(MAT_ID_ROW, x, "Mat ID:")
    ws.cell(
        NAME_ROW,
        x,
        f'=IF(OR({col}{TYPE_ROW}="",{col}{TYPE_ROW}="material"),_xlfn.XLOOKUP({col}{MAT_ID_ROW},mat!$A:$A, mat!$B:$B,"") & "","")',
    )
    ws.cell(SIZE_ROW, x, "Size:")
    ws.cell(XY_ROW, x, "x")
    ws.cell(XY_ROW, y, "y")
    # De keuzelijst van de cel Type, zoals bij de blokken van het sjabloon.
    for validation in ws.data_validations.dataValidation:
        if f"{get_column_letter(2)}{TYPE_ROW}" in validation.sqref:
            validation.add(f"{col}{TYPE_ROW}")


def write_xslope_zip(section: ProfileSlice, polygons: list[SoilPolygon], path: Path, label: str = "", library: dict | None = None) -> None:
    """Schrijft het XSLOPE-invoerbestand (zie write_xslope) en zijn stijlen (de
    kleuren van de grondsoorten, die XSLOPE ernaast bewaart) in een zip naar
    `path`, genoemd naar de zip: `<naam>.xlsx` en `<naam>_styles.json`, om
    samen uit te pakken. `label` wordt niet gebruikt (het staat in de bestandsnaam)."""
    with tempfile.TemporaryDirectory() as folder:
        xlsx = Path(folder) / f"{path.stem}.xlsx"
        write_xslope(section, polygons, xlsx, label, library)
        with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as archive:
            for file in sorted(Path(folder).iterdir()):
                archive.write(file, file.name)


def write_xslope(section: ProfileSlice, polygons: list[SoilPolygon], path: Path, label: str = "", library: dict | None = None) -> None:
    """Schrijft `polygons` (x = afstand tot de referentielijn, positief
    binnenwaarts; z = hoogte) als XSLOPE-invoerbestand naar `path`, en de
    kleuren van de materialen naar `<path stem>_styles.json` ernaast; met
    `library` (grondsoorten zoals dstability.read_stix_soils ze geeft, en
    `mapping`) krijgen de materialen de parameters van de grondsoorten die
    eraan gekoppeld zijn. `label` wordt niet gebruikt (het staat in de bestandsnaam)."""
    materials, material_of = _materials(section, polygons, library)
    path.with_name(f"{path.stem}_styles.json").write_text(json.dumps(_styles(materials), indent=2), encoding="utf-8")
    wb = openpyxl.load_workbook(TEMPLATE)

    # SI (het sjabloon staat op Imperial); geen tijdseenheid (geen tijdsafhankelijke grondwaterstroming).
    main = wb["main"]
    main["D8"] = "SI"
    main["D9"] = None
    main["D10"] = GAMMA_WATER

    mat = wb["mat"]
    header_row = next(
        r
        for r in range(1, 30)
        if str(mat.cell(r, 1).value).strip().lower() == "mat" and str(mat.cell(r, 2).value).strip().lower() == "name"
    )
    columns = {
        str(cell.value).strip().replace("_", ""): cell.column
        for cell in mat[header_row]
        if cell.value is not None and str(cell.value).strip()
    }
    for i, (name, props, _) in enumerate(materials):
        row = header_row + 1 + i
        mat.cell(row, columns["mat"], i + 1)
        mat.cell(row, columns["name"], name)
        for header, value in {**props, "u": "piezo"}.items():
            mat.cell(row, columns[header], value)

    # De geometrie zijn de polygonen: Max Depth van het blad 'profile' (de 0 van het sjabloon) wordt niet gebruikt.
    wb["profile"]["B2"] = None

    ws = wb["polygon"]
    for n, polygon in enumerate(polygons):
        x, y = 3 * n + 1, 3 * n + 2
        if ws.cell(HEADER_ROW, x).value is None:
            _add_block(ws, n)
        ws.cell(HEADER_ROW, x, f"Polygon #{n + 1}")
        ws.cell(TYPE_ROW, y, "material")
        ws.cell(MAT_ID_ROW, y, material_of[polygon.soil])
        for i, (px, pz) in enumerate(_points(polygon)):
            ws.cell(FIRST_VERTEX_ROW + i, x, px)
            ws.cell(FIRST_VERTEX_ROW + i, y, pz)

    # De materiaalnamen op het blad 'polygon' zijn formules: Excel rekent ze uit bij het openen.
    wb.calculation.fullCalcOnLoad = True
    wb.save(path)
